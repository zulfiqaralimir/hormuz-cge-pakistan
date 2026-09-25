import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

src='../data/raw/PBS_SUTs_and_IOT_2015-16.xlsx'
df=pd.read_excel(src,sheet_name='Input Output table',header=None)
codes=[str(df.iloc[r,0]).strip() for r in range(8,76)]
def sec(c):
    if c in ['01','02','03']: return 'Agri'
    if c in ['05-09','19','35']: return 'Energy'
    try:
        if 10<=int(c)<=33: return 'Mfg'
    except: pass
    return 'Serv'
m=pd.Series([sec(c) for c in codes],index=range(8,76))
S=['Agri','Energy','Mfg','Serv']
Z=df.iloc[8:76,2:70].astype(float); Z.index=range(8,76); Z.columns=range(8,76)
A=Z.groupby(m).sum().T.groupby(m).sum().T.loc[S,S]   # rows supplier, cols buyer
F=df.iloc[8:76,70:82].astype(float); F.index=range(8,76)
F.columns=['IC','HH','NPISH','GovC','GovI','GFCF','Val','Inv','ExpG','ExpS','Imp','Tot']
G=F.groupby(m).sum().loc[S]
def rowsum(r):
    s=df.iloc[r,2:70].astype(float); s.index=range(8,76); return s.groupby(m).sum()[S]
TAXA=rowsum(77); GVA=rowsum(78); GO=rowsum(79)
fdtax=df.iloc[77,71:80].astype(float).values  # HH,NPISH,GovC,GovI,GFCF,Val,Inv,ExpG,ExpS
tax_hh=fdtax[0]+fdtax[1]; tax_inv=fdtax[4]+fdtax[5]+fdtax[6]; tax_exp=fdtax[7]+fdtax[8]

LAB={'Agri':2274584,'Energy':69502,'Mfg':1420893,'Serv':4590194}
Wu=0.366*45283; Wr=0.634*30110   # HIES 2015-16 household weights
lab=lambda W,w,mx: W*(w+2/3*mx); cap=lambda W,mx,p,o: W*(mx/3+p+o)
Lu=lab(Wu,.5116,.2029); Lr=lab(Wr,.3205,.3987)
Ku=cap(Wu,.2029,.0303,.1416); Kr=cap(Wr,.3987,.0303,.0722)
u_lab=Lu/(Lu+Lr); u_cap=Ku/(Ku+Kr); print('shares',u_lab,u_cap)
u_cons=0.4665; u_inc=None
C=G.HH+G.NPISH; GC=G.GovC+G.GovI; I=G.GFCF+G.Val+G.Inv; X=G.ExpG+G.ExpS; M=G.Imp
DT=1191600; PEN=222500; DINT=1150800; FINT=112600; REM=2075664; u_rem=0.36

acc=['Act_Agri','Act_Energy','Act_Mfg','Act_Serv','Com_Agri','Com_Energy','Com_Mfg','Com_Serv',
     'Labour','Capital','HH_Urban','HH_Rural','Govt','RoW','S-I']
n=len(acc); idx={a:i for i,a in enumerate(acc)}
V={}  # (row,col)->value or formula marker
note={}
def put(r,c,v,nt=None):
    V[(r,c)]=float(v)
    if nt: note[(r,c)]=nt
labtot=sum(LAB.values()); captot=float(GVA.sum())-labtot
lab_u=labtot*u_lab; cap_u=captot*u_cap
inc_u=lab_u+cap_u; inc_share_u=inc_u/(labtot+captot)
for s in S:
    put('Com_'+s,'Act_'+s,GO[s],'Gross output, PBS IOT 2015-16')          # commodity pays activity
    for b in S: put('Com_'+s,'Act_'+b,0) if False else None
for s in S:
    for b in S: put('Com_'+s,'Act_'+b,A.loc[s,b],'Intermediate use, PBS IOT 2015-16')
# fix: activity receives GO from commodity -> row Act, col Com
V={k:v for k,v in V.items() if not (k[0].startswith('Com_') and k[1]=='Act_'+k[0][4:] and note.get(k,'').startswith('Gross'))}
for s in S:
    for b in S: put('Com_'+s,'Act_'+b,A.loc[s,b],'Intermediate use, PBS IOT 2015-16')
    put('Act_'+s,'Com_'+s,GO[s],'Gross output (basic prices), PBS IOT 2015-16')
    put('Labour','Act_'+s,LAB[s],'LFS 2014-15: employed x avg wage x 12')
    put('Capital','Act_'+s,GVA[s]-LAB[s],'GVA (IOT) minus labour')
    put('Govt','Act_'+s,TAXA[s],'Product taxes less subsidies on inputs, IOT')
    put('Com_'+s,'HH_Urban',C[s]*u_cons,'HH+NPISH consumption x urban share 46.65% (HIES)')
    put('Com_'+s,'HH_Rural',C[s]*(1-u_cons),'HH+NPISH consumption x rural share (HIES)')
    put('Com_'+s,'Govt',GC[s],'Govt final consumption, IOT')
    put('Com_'+s,'S-I',I[s],'GFCF+valuables+inventories, IOT')
    put('Com_'+s,'RoW',X[s],'Exports, IOT')
    put('RoW','Com_'+s,M[s],'Imports, IOT')
put('HH_Urban','Labour',lab_u,'Urban share = wages + 2/3 of mixed income (HIES Table 11)')
put('HH_Rural','Labour',labtot-lab_u)
put('HH_Urban','Capital',cap_u,'Urban share = 1/3 mixed income + property + owner-occupied housing (HIES)')
put('HH_Rural','Capital',captot-cap_u)
put('HH_Urban','Govt',(PEN+DINT)*inc_share_u,'Pensions 222.5bn + domestic interest 1,150.8bn (MoF FY16) x income share')
put('HH_Rural','Govt',(PEN+DINT)*(1-inc_share_u))
put('HH_Urban','RoW',REM*u_rem,'Remittances $19.92bn x 104.2 (SBP), 36% urban (HIES)')
put('HH_Rural','RoW',REM*(1-u_rem))
put('Govt','HH_Urban',DT*inc_share_u+tax_hh*u_cons,'Direct taxes (MoF) x income share + product taxes on consumption (IOT)')
put('Govt','HH_Rural',DT*(1-inc_share_u)+tax_hh*(1-u_cons))
put('Govt','S-I',tax_inv,'Product taxes on investment, IOT')
put('Govt','RoW',tax_exp,'Product taxes on exports, IOT')
put('RoW','Govt',FINT,'Foreign interest (MoF FY16)')
resid={('S-I','HH_Urban'):'HH_Urban',('S-I','HH_Rural'):'HH_Rural',('S-I','Govt'):'Govt',('S-I','RoW'):'RoW'}

wb=Workbook(); ws=wb.active; ws.title='SAM'
f=Font(name='Arial',size=9); fb=Font(name='Arial',size=9,bold=True); fblue=Font(name='Arial',size=9,color='0000FF')
hdr=PatternFill('solid',fgColor='DDEBF7'); yel=PatternFill('solid',fgColor='FFFF00'); thin=Side(style='thin',color='BBBBBB')
ws['A1']='Pakistan SAM 2015-16 (PKR million) — rows = receipts, columns = payments'; ws['A1'].font=Font(name='Arial',size=11,bold=True)
R0=3; C0=2
ws.cell(R0,1,'').fill=hdr
for j,a in enumerate(acc):
    c=ws.cell(R0,C0+j,a); c.font=fb; c.fill=hdr; c.alignment=Alignment(text_rotation=90,horizontal='center')
    r=ws.cell(R0+1+j,1,a); r.font=fb; r.fill=hdr
ws.cell(R0,C0+n,'TOTAL').font=fb; ws.cell(R0,C0+n).fill=hdr
tr=R0+1+n; ws.cell(tr,1,'TOTAL').font=fb; ws.cell(tr,1).fill=hdr
ws.cell(tr+1,1,'Row − Col (check)').font=fb
for i,a in enumerate(acc):
    for j,b in enumerate(acc):
        cell=ws.cell(R0+1+i,C0+j)
        if (a,b) in V:
            cell.value=round(V[(a,b)],1); cell.font=fblue
            if (a,b) in note: cell.comment=Comment(note[(a,b)],'SAM')
        cell.number_format='#,##0;(#,##0);-'; cell.border=Border(top=thin,bottom=thin,left=thin,right=thin)
        if cell.font!=fblue: cell.font=f
# residual savings: column total must equal row total
for (a,b),acct in resid.items():
    i=idx[a]; j=idx[b]; k=idx[acct]
    col=L(C0+j); rowtot=f'{L(C0+n)}{R0+1+k}'
    others=f'SUM({col}{R0+1}:{col}{R0+n-1})'
    cell=ws.cell(R0+1+i,C0+j,f'={rowtot}-({others})'); cell.font=f; cell.fill=yel
    cell.comment=Comment('Savings = residual (account income − other spending)','SAM')
for i in range(n):
    rr=R0+1+i; c=ws.cell(rr,C0+n,f'=SUM({L(C0)}{rr}:{L(C0+n-1)}{rr})'); c.font=fb; c.number_format='#,##0'
for j in range(n):
    col=L(C0+j); c=ws.cell(tr,C0+j,f'=SUM({col}{R0+1}:{col}{R0+n})'); c.font=fb; c.number_format='#,##0'
    d=ws.cell(tr+1,C0+j,f'={L(C0+n)}{R0+1+j}-{col}{tr}'); d.font=fb; d.number_format='#,##0;(#,##0);0'
ws.column_dimensions['A'].width=18
for j in range(n+1): ws.column_dimensions[L(C0+j)].width=12
ws.row_dimensions[R0].height=70
nr=tr+3
for t in ['Blue = data (hover for source). Yellow = savings residuals (formulas).',
          'Sources: PBS SUT/IOT 2015-16; LFS 2014-15; HIES 2015-16; MoF/SBP fiscal & BoP FY16.',
          'Energy = PSIC 05-09 mining, 19 refining, 35 electricity/gas. Govt includes the tax account.']:
    ws.cell(nr,1,t).font=f; nr+=1
ws.freeze_panes='B4'
wb.save('Pakistan_SAM_2015-16.xlsx')
